"""
DOCX document parser (requires python-docx).

Extracts paragraphs with their heading style levels to reconstruct
a section hierarchy similar to the Markdown parser output.
"""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import List, Optional

from eryc.ingestion.parsers.base import ParsedDocument, ParsedSection, ParserBase


class DocxParser(ParserBase):
    """Parse DOCX files into hierarchical sections via python-docx."""

    supported_extensions = [".docx"]

    def parse(self, path: Path, *, doc_type: Optional[str] = None) -> ParsedDocument:
        from docx import Document  # type: ignore

        doc = Document(str(path))
        sections: List[ParsedSection] = []
        path_stack: List[tuple[int, str]] = []
        current_heading = ""
        current_level = 0
        current_lines: List[str] = []

        def _flush() -> None:
            if not current_lines:
                return
            sp = "/".join(h for _, h in path_stack) or "Root"
            sections.append(
                ParsedSection(
                    section_path=sp,
                    heading=current_heading,
                    text="\n".join(current_lines).strip(),
                    level=current_level,
                )
            )

        for para in doc.paragraphs:
            style_name = para.style.name or ""
            text = para.text.strip()
            if not text:
                continue

            if style_name.startswith("Heading"):
                _flush()
                current_lines = []
                try:
                    level = int(style_name.split()[-1])
                except ValueError:
                    level = 1
                path_stack = [(l, h) for l, h in path_stack if l < level]
                path_stack.append((level, text))
                current_heading = text
                current_level = level
            else:
                current_lines.append(text)

        _flush()

        raw_text = "\n\n".join(s.text for s in sections)
        content_hash = hashlib.sha256(raw_text.encode()).hexdigest()

        # Derive title from first Heading 1 or filename.
        title = path.stem
        for s in sections:
            if s.level == 1:
                title = s.heading
                break

        return ParsedDocument(
            source_path=str(path),
            canonical_title=title,
            sections=sections or [
                ParsedSection(section_path="Root", heading="", text=raw_text, level=0)
            ],
            raw_text=raw_text,
            doc_type=doc_type or "docx",
            content_hash=content_hash,
            metadata={"file_size": path.stat().st_size},
        )
