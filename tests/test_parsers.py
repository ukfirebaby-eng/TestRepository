import csv as _csv
import uuid
import pytest
from unittest.mock import MagicMock

from core.orchestrator import DiamondOrchestrator


def _make_orchestrator():
    return DiamondOrchestrator(
        tenant_id="t1",
        document_id=f"doc_{uuid.uuid4().hex[:8]}",
        vault=MagicMock(),
    )


def _assert_chunks(chunks):
    assert isinstance(chunks, list)
    assert len(chunks) >= 1
    for c in chunks:
        assert "chunk_id" in c
        assert "text" in c
        assert "page" in c
        assert "bbox" in c
        assert len(c["text"]) > 20


class TestParsePlaintext:
    def test_txt_returns_chunks(self, tmp_path):
        f = tmp_path / "doc.txt"
        f.write_text(
            "First paragraph with enough content to exceed twenty characters.\n\n"
            "Second paragraph also has enough text to pass the filter threshold."
        )
        chunks = _make_orchestrator()._parse_plaintext(str(f))
        _assert_chunks(chunks)

    def test_md_strips_markdown_markers(self, tmp_path):
        f = tmp_path / "doc.md"
        f.write_text(
            "# Heading One\n\n"
            "**Bold text** and _italic_ content that is definitely long enough.\n\n"
            "Another paragraph with `code` markers and enough words to pass."
        )
        chunks = _make_orchestrator()._parse_plaintext(str(f))
        _assert_chunks(chunks)
        for c in chunks:
            assert "#" not in c["text"]
            assert "**" not in c["text"]

    def test_short_paragraphs_are_filtered(self, tmp_path):
        f = tmp_path / "short.txt"
        f.write_text("Hi\n\nThis paragraph is definitely long enough to survive the filter.")
        chunks = _make_orchestrator()._parse_plaintext(str(f))
        assert all(len(c["text"]) > 20 for c in chunks)

    def test_page_and_bbox_are_sentinels(self, tmp_path):
        f = tmp_path / "doc.txt"
        f.write_text("A sufficiently long paragraph to make it past the twenty character filter.")
        chunks = _make_orchestrator()._parse_plaintext(str(f))
        for c in chunks:
            assert c["page"] == 0
            assert c["bbox"] is None


class TestParseDocx:
    def test_docx_returns_chunks(self, tmp_path):
        from docx import Document
        doc = Document()
        doc.add_paragraph("This is the first paragraph with plenty of content to analyse.")
        doc.add_paragraph("This is the second paragraph also with plenty of content to analyse.")
        path = tmp_path / "doc.docx"
        doc.save(str(path))

        chunks = _make_orchestrator()._parse_docx(str(path))
        _assert_chunks(chunks)

    def test_empty_paragraphs_are_filtered(self, tmp_path):
        from docx import Document
        doc = Document()
        doc.add_paragraph("")
        doc.add_paragraph("   ")
        doc.add_paragraph("This paragraph has enough text to survive the length filter threshold.")
        path = tmp_path / "sparse.docx"
        doc.save(str(path))

        chunks = _make_orchestrator()._parse_docx(str(path))
        assert len(chunks) == 1

    def test_page_and_bbox_are_sentinels(self, tmp_path):
        from docx import Document
        doc = Document()
        doc.add_paragraph("A paragraph with sufficient length to pass the twenty-character filter.")
        path = tmp_path / "doc.docx"
        doc.save(str(path))

        chunks = _make_orchestrator()._parse_docx(str(path))
        for c in chunks:
            assert c["page"] == 0
            assert c["bbox"] is None


class TestParseTabular:
    def test_csv_returns_chunks(self, tmp_path):
        path = tmp_path / "data.csv"
        with open(str(path), "w", newline="") as f:
            writer = _csv.writer(f)
            writer.writerow(["Project", "Owner", "Status", "Deadline", "Budget"])
            for i in range(15):
                writer.writerow([f"Project {i}", f"Owner {i}", "Active", "2025-12-01", f"£{i * 1000}"])

        chunks = _make_orchestrator()._parse_tabular(str(path))
        _assert_chunks(chunks)
        # 15 data rows + 1 header = 16 rows → 2 chunks of 10
        assert len(chunks) == 2

    def test_xlsx_returns_chunks(self, tmp_path):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Risk", "Likelihood", "Impact", "Owner", "Mitigation"])
        for i in range(12):
            ws.append([f"Risk {i}", "High", "Critical", f"Team {i}", f"Action {i}"])
        path = tmp_path / "data.xlsx"
        wb.save(str(path))

        chunks = _make_orchestrator()._parse_tabular(str(path))
        _assert_chunks(chunks)

    def test_page_and_bbox_are_sentinels(self, tmp_path):
        path = tmp_path / "data.csv"
        with open(str(path), "w", newline="") as f:
            writer = _csv.writer(f)
            for i in range(5):
                writer.writerow([f"Value {i}", f"Label {i}", f"Category {i}", f"Note {i}", f"Tag {i}"])

        chunks = _make_orchestrator()._parse_tabular(str(path))
        for c in chunks:
            assert c["page"] == 0
            assert c["bbox"] is None


class TestParsePptx:
    def test_pptx_returns_chunks(self, tmp_path):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        layout = prs.slide_layouts[1]
        for i in range(3):
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = f"Slide {i + 1} Title"
            slide.placeholders[1].text = (
                f"This is the body text for slide {i + 1} with enough content to pass the filter."
            )
        path = tmp_path / "deck.pptx"
        prs.save(str(path))

        chunks = _make_orchestrator()._parse_pptx(str(path))
        _assert_chunks(chunks)
        assert len(chunks) == 3

    def test_slide_index_used_as_page(self, tmp_path):
        from pptx import Presentation
        prs = Presentation()
        layout = prs.slide_layouts[1]
        for i in range(2):
            slide = prs.slides.add_slide(layout)
            slide.shapes.title.text = f"Slide {i + 1}"
            slide.placeholders[1].text = f"Body content for slide {i + 1} is long enough here."
        path = tmp_path / "deck.pptx"
        prs.save(str(path))

        chunks = _make_orchestrator()._parse_pptx(str(path))
        pages = [c["page"] for c in chunks]
        assert pages == [1, 2]

    def test_bbox_is_none(self, tmp_path):
        from pptx import Presentation
        prs = Presentation()
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = "Title"
        slide.placeholders[1].text = "Body text that is definitely long enough to survive the filter."
        path = tmp_path / "deck.pptx"
        prs.save(str(path))

        chunks = _make_orchestrator()._parse_pptx(str(path))
        for c in chunks:
            assert c["bbox"] is None


class TestParseDocumentRouter:
    def test_routes_pdf_to_pdf_parser(self, tmp_path, monkeypatch):
        orch = _make_orchestrator()
        called = {}
        monkeypatch.setattr(orch, "_parse_pdf_with_geometry", lambda p: called.update({"pdf": p}) or [])
        orch._parse_document(str(tmp_path / "report.pdf"))
        assert "pdf" in called

    def test_routes_txt_to_plaintext_parser(self, tmp_path, monkeypatch):
        f = tmp_path / "doc.txt"
        f.write_text("content")
        orch = _make_orchestrator()
        called = {}
        monkeypatch.setattr(orch, "_parse_plaintext", lambda p: called.update({"txt": p}) or [])
        orch._parse_document(str(f))
        assert "txt" in called

    def test_routes_md_to_plaintext_parser(self, tmp_path, monkeypatch):
        f = tmp_path / "doc.md"
        f.write_text("content")
        orch = _make_orchestrator()
        called = {}
        monkeypatch.setattr(orch, "_parse_plaintext", lambda p: called.update({"md": p}) or [])
        orch._parse_document(str(f))
        assert "md" in called

    def test_raises_for_unsupported_extension(self, tmp_path):
        orch = _make_orchestrator()
        with pytest.raises(ValueError, match="Unsupported file type"):
            orch._parse_document(str(tmp_path / "virus.exe"))
