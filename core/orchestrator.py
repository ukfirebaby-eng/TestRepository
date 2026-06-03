import fitz  # PyMuPDF
import uuid
import re
import csv
import os
from datetime import datetime, timezone
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import List, Dict, Any

# Import our previously written modules
from core.vault import HybridVault
from core.agents import DeconstructorAgent, ContradictionHunterAgent, FragilityAgent, ChronosAgent, _get_model
from core.accuracy import claim_extractor
from core.accuracy.deterministic_validator import validate_claim
from core.accuracy.entity_canonicalizer import build_canonical_entities_from_claims
from core.accuracy.evidence_span_builder import build_evidence_spans, compute_source_hash
from core.accuracy.graph_promoter import promote_claim_to_topology
from core.accuracy.schemas import DocumentManifest
from core.risk_finding import build_risk_finding

HUB_MIN_DEPENDENTS = 3


class DiamondOrchestrator:
    def __init__(self, tenant_id: str, document_id: str, document_name: str = "Untitled", vault: HybridVault = None, log_fn=None):
        self.tenant_id = tenant_id
        self.document_id = document_id
        self.document_name = document_name
        self.vault = vault if vault is not None else HybridVault(tenant_id=tenant_id)
        self._log_fn = log_fn
        self.accuracy_metrics: Dict[str, int] = {}

    def _emit(self, msg: str) -> None:
        """Prints to console and forwards to the optional UI log callback."""
        print(msg)
        if self._log_fn:
            self._log_fn(msg)

    def _claim_layer_enabled(self) -> bool:
        return os.environ.get("DIAMOND_MINER_CLAIM_LAYER", "").strip() == "1"

    def _claim_promotion_enabled(self) -> bool:
        return os.environ.get("DIAMOND_MINER_USE_CLAIM_PROMOTION", "").strip() == "1"

    def _parse_pdf_with_geometry(self, file_path: str) -> List[Dict[str, Any]]:
        """
        Act I: Reads the PDF and extracts block-level text alongside (x0, y0, x1, y1) coordinates.
        """
        doc = fitz.open(file_path)
        chunks = []

        for page_num, page in enumerate(doc):
            blocks = page.get_text("dict").get("blocks", [])
            for block in blocks:
                if block.get("type") == 0:  # Type 0 is text
                    chunk_text = " ".join([span.get("text", "")
                                           for line in block.get("lines", [])
                                           for span in line.get("spans", [])]).strip()

                    if len(chunk_text) > 20:  # Filter out tiny artifacts
                        chunks.append({
                            "chunk_id": str(uuid.uuid4()),
                            "text": chunk_text,
                            "page": page_num + 1,
                            "bbox": block.get("bbox")  # (x0, y0, x1, y1)
                        })
        return chunks

    def _parse_document(self, file_path: str) -> List[Dict[str, Any]]:
        """Routes to the appropriate parser based on file extension."""
        ext = Path(file_path).suffix.lower()
        if ext == ".pdf":
            return self._parse_pdf_with_geometry(file_path)
        elif ext in (".txt", ".md"):
            return self._parse_plaintext(file_path)
        elif ext == ".docx":
            return self._parse_docx(file_path)
        elif ext in (".csv", ".xlsx"):
            return self._parse_tabular(file_path)
        elif ext == ".pptx":
            return self._parse_pptx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

    def _parse_plaintext(self, file_path: str) -> List[Dict[str, Any]]:
        """Parses .txt and .md files by splitting on blank lines."""
        with open(file_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        # Strip Markdown syntax markers
        content = re.sub(r"[#*_`]", "", content)
        raw_chunks = content.split("\n\n")
        chunks = []
        for raw in raw_chunks:
            text = raw.strip()
            if len(text) > 20:
                chunks.append({
                    "chunk_id": str(uuid.uuid4()),
                    "text": text,
                    "page": 0,
                    "bbox": None,
                })
        return chunks

    def _parse_docx(self, file_path: str) -> List[Dict[str, Any]]:
        """Parses .docx files by iterating paragraphs."""
        from docx import Document
        doc = Document(file_path)
        chunks = []
        for para in doc.paragraphs:
            text = para.text.strip()
            if len(text) > 20:
                chunks.append({
                    "chunk_id": str(uuid.uuid4()),
                    "text": text,
                    "page": 0,
                    "bbox": None,
                })
        return chunks

    def _parse_tabular(self, file_path: str) -> List[Dict[str, Any]]:
        """Parses .csv and .xlsx files by grouping rows into chunks of 10."""
        ext = Path(file_path).suffix.lower()
        rows = []
        if ext == ".csv":
            with open(file_path, "r", encoding="utf-8", errors="replace", newline="") as f:
                reader = csv.reader(f)
                for row in reader:
                    joined = " | ".join(cell.strip() for cell in row if cell.strip())
                    if joined:
                        rows.append(joined)
        else:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    joined = " | ".join(str(c).strip() for c in row if c is not None and str(c).strip())
                    if joined:
                        rows.append(joined)
            wb.close()

        chunks = []
        for i in range(0, len(rows), 10):
            text = "\n".join(rows[i:i + 10])
            if len(text.split()) >= 5:
                chunks.append({
                    "chunk_id": str(uuid.uuid4()),
                    "text": text,
                    "page": 0,
                    "bbox": None,
                })
        return chunks

    def _parse_pptx(self, file_path: str) -> List[Dict[str, Any]]:
        """Parses .pptx files; each slide maps to a page number."""
        from pptx import Presentation
        prs = Presentation(file_path)
        chunks = []
        for slide_index, slide in enumerate(prs.slides):
            parts = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    parts.append(shape.text_frame.text.strip())
            text = "\n".join(p for p in parts if p)
            if len(text) > 20:
                chunks.append({
                    "chunk_id": str(uuid.uuid4()),
                    "text": text,
                    "page": slide_index + 1,
                    "bbox": None,
                })
        return chunks

    def _process_single_chunk(self, chunk: Dict[str, Any]) -> None:
        """
        Act II: The worker function for a single thread.
        Saves semantic text, builds graph topology, and extracts temporal metadata.
        ChromaDB insert and topology extraction are fired concurrently since they
        are independent — this removes one full serial step per chunk.
        """
        with ThreadPoolExecutor(max_workers=2) as inner:
            embed_future = inner.submit(
                self.vault.insert_document_chunk,
                chunk_id=chunk["chunk_id"],
                document_id=self.document_id,
                text=chunk["text"],
                page=chunk["page"],
                bbox=chunk["bbox"],
            )
            topo_future = inner.submit(DeconstructorAgent.extract_topology, chunk["text"])
            embed_future.result()
            topology = topo_future.result()

        if topology.get("nodes") and topology.get("edges"):
            self.vault.insert_graph_topology(
                nodes=topology["nodes"],
                edges=topology["edges"],
                source_chunk_id=chunk["chunk_id"],
                document_id=self.document_id
            )

            # 3. Temporal Pass: extract ISO 8601 dates for the nodes just created
            node_ids = [node["id"] for node in topology["nodes"]]
            temporal_data = ChronosAgent.extract_time_data(chunk["text"], existing_nodes=node_ids)

            if temporal_data.get("temporal_nodes"):
                self.vault.insert_temporal_data(temporal_data["temporal_nodes"])

            # Store temporal edges (STARTS_AFTER) in the existing edges table
            if temporal_data.get("temporal_edges"):
                temporal_edge_dicts = [
                    {
                        "source_id": te["source_id"],
                        "target_id": te["target_id"],
                        "relationship": te["relationship"],
                    }
                    for te in temporal_data["temporal_edges"]
                ]
                self.vault.insert_graph_topology(
                    nodes=[],
                    edges=temporal_edge_dicts,
                    source_chunk_id=chunk["chunk_id"],
                    document_id=self.document_id
                )

    def run_ingestion_pipeline(self, file_path: str, max_workers: int = 10) -> None:
        """
        The Main Execution Loop. Orchestrates parallel processing and enforces the synchronization barrier.
        """
        self.vault.insert_document(self.document_id, self.document_name)
        self._emit(f"[*] Orchestrator: Parsing document {self.document_id}...")
        chunks = self._parse_document(file_path)
        source_text = "\n\n".join(chunk["text"] for chunk in chunks)
        source_hash = compute_source_hash(source_text)
        if self._claim_layer_enabled():
            manifest = DocumentManifest(
                document_id=self.document_id,
                filename=self.document_name,
                source_hash=source_hash,
                ingested_at=datetime.now(timezone.utc),
                llm_model=_get_model("fast"),
            )
            self.vault.save_document_manifest(manifest)
            spans = build_evidence_spans(
                document_id=self.document_id,
                chunks=chunks,
                source_hash=source_hash,
            )
            self.vault.insert_evidence_spans(spans)
            extraction_result = claim_extractor.extract_claims_with_failures(spans)
            claims = extraction_result.claims
            self.vault.insert_extraction_failures(extraction_result.failures)
            known_span_ids = {span.span_id for span in spans}
            validation_results = [
                validate_claim(claim, known_span_ids=known_span_ids)
                for claim in claims
            ]
            validated_claims = [
                claim.model_copy(update={"validation_status": result.status})
                for claim, result in zip(claims, validation_results)
            ]
            self.vault.insert_extracted_claims(validated_claims)
            self.vault.insert_validation_results(validation_results)
            canonical_entities = build_canonical_entities_from_claims(validated_claims)
            self.vault.insert_canonical_entities(self.document_id, canonical_entities)
            self.accuracy_metrics = {
                "evidence_spans": len(spans),
                "claims": len(claims),
                "canonical_entities": len(canonical_entities),
                "validated": sum(1 for result in validation_results if result.status == "passed"),
                "needs_review": sum(1 for result in validation_results if result.status == "needs_review"),
                "failed": sum(1 for result in validation_results if result.status == "failed"),
                "promotable": sum(1 for result in validation_results if result.can_promote),
                "extraction_failures": len(extraction_result.failures),
            }
            self.accuracy_metrics.update(extraction_result.metrics)
            if self._claim_promotion_enabled():
                promoted_nodes = []
                promoted_edges = []
                for claim in validated_claims:
                    topology = promote_claim_to_topology(claim)
                    promoted_nodes.extend(topology["nodes"])
                    promoted_edges.extend(topology["edges"])
                if promoted_nodes or promoted_edges:
                    self.vault.insert_graph_topology(
                        nodes=promoted_nodes,
                        edges=promoted_edges,
                        source_chunk_id="claim_layer",
                        document_id=self.document_id,
                    )
            promotable_count = sum(1 for result in validation_results if result.can_promote)
            self._emit(
                f"[*] Claim Layer: Stored {len(spans)} evidence span(s), "
                f"{len(claims)} claim(s), {promotable_count} promotable, "
                f"{len(extraction_result.failures)} extraction failure(s)."
            )
        self._emit(f"[*] Orchestrator: Extracted {len(chunks)} geometric chunks. Beginning parallel Deconstruction.")

        # Spin up concurrent threads to blast through the document chunk-by-chunk
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            futures = [executor.submit(self._process_single_chunk, chunk) for chunk in chunks]

            # The Synchronization Barrier: Wait for all threads to finish
            for count, future in enumerate(as_completed(futures), 1):
                try:
                    future.result()  # This will raise any exceptions caught in the thread
                    if count % 10 == 0:
                        self._emit(f"    -> Processed {count}/{len(chunks)} chunks...")
                except Exception as e:
                    self._emit(f"[!] Orchestrator Thread Failed: {e}")

        self._emit("[*] Orchestrator: Ingestion complete. Knowledge Graph is locked and loaded.")

    def interrogate_friction(self) -> List[Dict[str, Any]]:
        """
        Act III: The Hunt. Queries SQLite for paradoxes and asks the AI to synthesize mitigations.
        """
        # Zero-cost structural query
        conflicts = self.vault.get_triangular_conflicts(document_id=self.document_id)
        self._emit(f"[*] Contradiction Hunter: Found {len(conflicts)} structural conflict(s). Synthesizing mitigations...")

        CONFIDENCE_THRESHOLD = 0.65  # Discard weak or spurious conflicts

        verified_diamonds = []
        spurious_count = 0

        def _verify_conflict(args):
            i, conflict = args
            self._emit(f"    -> Verifying conflict {i}/{len(conflicts)}...")
            prov_requires = self.vault.get_chunk_provenance(conflict["chunk_requires"])
            prov_blocks = self.vault.get_chunk_provenance(conflict["chunk_blocks"])
            combined_text = f"Claim 1: {prov_requires.get('text')}\nClaim 2: {prov_blocks.get('text')}"
            names = self.vault.get_node_names([conflict['node_a'], conflict['node_b'], conflict['node_c']])
            name_a = names.get(conflict['node_a'], conflict['node_a'])
            name_b = names.get(conflict['node_b'], conflict['node_b'])
            name_c = names.get(conflict['node_c'], conflict['node_c'])
            structural_clash = f'"{name_a}" REQUIRES "{name_b}", but "{name_c}" BLOCKS "{name_b}"'
            result = ContradictionHunterAgent.synthesize_mitigation(structural_clash, combined_text)
            result_for_finding = dict(result)
            raw_finding = result.get("finding") if isinstance(result.get("finding"), dict) else {}
            claim_ids = conflict.get("claim_ids") or []
            evidence_span_ids = conflict.get("evidence_span_ids") or []
            if claim_ids or evidence_span_ids:
                result_for_finding["finding"] = {
                    **raw_finding,
                    "claim_ids": claim_ids,
                    "evidence_span_ids": evidence_span_ids,
                    "claim_validation_status": raw_finding.get("claim_validation_status") or "passed",
                }
            finding = build_risk_finding(
                result_for_finding,
                risk_type="structural",
                source_name=name_a,
                target_name=name_c,
                relationship="structural friction",
                blocked_work=name_a,
                blocking_condition=name_b,
            )
            return conflict, result, finding

        with ThreadPoolExecutor(max_workers=3) as executor:
            for conflict, result, finding in executor.map(_verify_conflict, enumerate(conflicts, 1)):
                if not result.get("is_genuine") or result.get("confidence", 0) < CONFIDENCE_THRESHOLD:
                    spurious_count += 1
                    self._emit(f"       [skipped — spurious] confidence={result.get('confidence', 0):.2f}")
                    continue
                verified_diamonds.append({
                    "source": conflict["node_a"],
                    "target": conflict["node_c"],
                    "diamond": result["analysis"],
                    "provenance_ids": [conflict["chunk_requires"], conflict["chunk_blocks"]],
                    "severity": result.get("severity", 3),
                    "probability": result.get("probability", 3),
                    "finding": finding,
                })

        self._emit(f"[*] Contradiction Hunter: {len(verified_diamonds)} genuine conflict(s) confirmed, {spurious_count} spurious patterns discarded.")
        self.vault.upsert_friction_lines(self.document_id, verified_diamonds)
        return verified_diamonds

    def interrogate_fragility(self) -> None:
        """
        DLI Pass: Identifies hub nodes and verifies their cascade fragility.
        """
        self._emit("[*] DLI: Mapping fragility...")

        CONFIDENCE_THRESHOLD = 0.65

        hub_nodes = self.vault.get_hub_nodes(self.document_id, min_dependents=HUB_MIN_DEPENDENTS)
        self._emit(f"[*] DLI: Found {len(hub_nodes)} hub node candidate(s). Verifying...")

        verified = []
        spurious_count = 0

        def _verify_hub(args):
            i, hub = args
            self._emit(f"    -> Verifying hub {i}/{len(hub_nodes)}: {hub['name']}...")
            dependent_names = list(self.vault.get_node_names(hub["dependent_node_ids"]).values())
            source_chunk_id = self.vault.get_node_source_chunk(hub["id"], self.document_id)
            if source_chunk_id is None:
                self._emit(f"       [skipped — no source chunk found for hub node {hub['id']}]")
                return hub, None
            chunk_provenance = self.vault.get_chunk_provenance(source_chunk_id)
            source_text = chunk_provenance.get("text", "")
            result = FragilityAgent.analyse(hub["name"], dependent_names, source_text)
            return hub, result

        with ThreadPoolExecutor(max_workers=3) as executor:
            for hub, result in executor.map(_verify_hub, enumerate(hub_nodes, 1)):
                if result is None:
                    continue
                if not result.get("is_genuine") or result.get("confidence", 0) < CONFIDENCE_THRESHOLD:
                    spurious_count += 1
                    self._emit(f"       [skipped — spurious] confidence={result.get('confidence', 0):.2f}")
                    continue
                verified.append({
                    "hub_node_id": hub["id"],
                    "dependency_count": hub["dependency_count"],
                    "insight": result["insight"],
                    "cascade_nodes": result["cascade_nodes"]
                })

        self.vault.upsert_fragility_lines(self.document_id, verified)
        self._emit(f"[*] DLI: {len(verified)} fragility point(s) confirmed, {spurious_count} spurious patterns discarded.")

    def interrogate_time_friction(self) -> List[Dict[str, Any]]:
        """
        Act IV: The Timeline Hunt. Queries SQLite for Negative Slack and
        synthesizes chronological mitigation strategies.
        """
        conflicts = self.vault.get_chronological_friction(document_id=self.document_id)
        self._emit(f"[*] Chronos: Found {len(conflicts)} chronological conflict(s). Synthesizing mitigations...")

        CONFIDENCE_THRESHOLD = 0.65
        verified_time_diamonds = []
        spurious_count = 0

        def _verify_time_conflict(args):
            i, conflict = args
            self._emit(f"    -> Verifying time conflict {i}/{len(conflicts)}...")
            prov_chunk = self.vault.get_chunk_provenance(conflict["chunk_bridge"])
            names = self.vault.get_node_names([conflict["predecessor"], conflict["successor"]])
            pred_name = names.get(conflict["predecessor"], conflict["predecessor"])
            succ_name = names.get(conflict["successor"], conflict["successor"])
            clash_payload = (
                f"Node '{succ_name}' is scheduled to start on {conflict['succ_start']}. "
                f"However, it MUST START AFTER Node '{pred_name}', "
                f"which does not end until {conflict['pred_end']}."
            )
            result = ContradictionHunterAgent.synthesize_mitigation(
                structural_clash=clash_payload,
                source_text=prov_chunk.get("text", "No source text found.")
            )
            finding = build_risk_finding(
                result,
                risk_type="timeline",
                source_name=pred_name,
                target_name=succ_name,
                relationship="chronological friction",
                blocked_work=succ_name,
                blocking_condition=pred_name,
            )
            return conflict, result, finding

        with ThreadPoolExecutor(max_workers=3) as executor:
            for conflict, result, finding in executor.map(_verify_time_conflict, enumerate(conflicts, 1)):
                if not result.get("is_genuine") or result.get("confidence", 0) < CONFIDENCE_THRESHOLD:
                    spurious_count += 1
                    self._emit(f"       [skipped — spurious] confidence={result.get('confidence', 0):.2f}")
                    continue
                verified_time_diamonds.append({
                    "type": "chronological",
                    "source": conflict["predecessor"],
                    "target": conflict["successor"],
                    "diamond": result["analysis"],
                    "provenance_ids": [conflict["chunk_bridge"]],
                    "severity": result.get("severity", 3),
                    "probability": result.get("probability", 3),
                    "finding": finding,
                })

        self._emit(f"[*] Chronos: {len(verified_time_diamonds)} genuine time conflict(s), {spurious_count} spurious discarded.")
        self.vault.upsert_chronological_friction_lines(self.document_id, verified_time_diamonds)
        return verified_time_diamonds
